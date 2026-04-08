#!/usr/bin/env python3
"""CSV and status generation for auto-job-hunt."""

from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path

from paths import jobs_dir, jobs_status_dir, sanitize_filename


def jobs_to_csv(jobs: list[dict], output_path: str | Path) -> str:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    if str(path).endswith(".xlsx"):
        return jobs_to_xlsx(jobs, path)

    with path.open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(get_csv_headers())
        for job in jobs:
            writer.writerow(extract_job_row(job))

    return str(path)


def jobs_to_xlsx(jobs: list[dict], output_path: Path) -> str:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError(
            "XLSX output requires openpyxl. Install openpyxl to generate .xlsx files."
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"

    headers = get_csv_headers()
    header_fill = PatternFill(
        start_color="366092", end_color="366092", fill_type="solid"
    )
    header_font = Font(bold=True, color="FFFFFF")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, job in enumerate(jobs, 2):
        for col_idx, value in enumerate(extract_job_row(job), 1):
            if value is None:
                value = ""
            elif isinstance(value, bool):
                value = "Yes" if value else "No"
            elif isinstance(value, (list, set)):
                value = ", ".join(str(v) for v in value)
            ws.cell(row=row_idx, column=col_idx, value=value)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 15

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 50
    ws.column_dimensions["P"].width = 40

    wb.save(output_path)
    return str(output_path)


def get_csv_headers() -> list[str]:
    return [
        "id",
        "externalJobId",
        "title",
        "jobDescription",
        "salaryMin",
        "salaryMax",
        "salaryCurrency",
        "salaryType",
        "jobTypes",
        "workType",
        "experienceLevel",
        "visaSponsorship",
        "city",
        "state",
        "country",
        "remote",
        "postedDate",
        "applyUrl",
        "easyApply",
        "companyName",
        "companyIndustry",
        "companyWebsite",
        "companyRating",
        "jobSource",
        "score",
        "recommendation",
        "searchTerm",
        "source",
        "username",
        "role",
        "profilePath",
        "resumePath",
    ]


def extract_job_row(job: dict) -> list:
    salary = job.get("salary", {})
    location = job.get("location", {})
    company = job.get("company", {})
    posting = job.get("posting", {})
    apply = job.get("apply", {})
    job_source = job.get("jobSource", {})

    def safe_get(d: dict, key: str, default=""):
        val = d.get(key, default)
        return default if val is None else val

    return [
        safe_get(job, "id"),
        safe_get(job, "externalJobId"),
        safe_get(job, "title"),
        (safe_get(job, "jobDescription") or "")[:500],
        safe_get(salary, "min"),
        safe_get(salary, "max"),
        safe_get(salary, "currencyCode", "USD"),
        safe_get(salary, "type"),
        ", ".join(job.get("jobTypes", [])),
        safe_get(job, "workType"),
        safe_get(job, "experienceLevel"),
        job.get("visaSponsorship", False),
        safe_get(location, "city"),
        safe_get(location, "state"),
        safe_get(location, "country"),
        location.get("remote", False),
        safe_get(posting, "postedDate"),
        safe_get(apply, "applyUrl"),
        apply.get("easyApply", False),
        safe_get(company, "name"),
        safe_get(company, "industry"),
        safe_get(company, "website"),
        company.get("rating"),
        safe_get(job_source, "sourceName"),
        job.get("score", 0),
        safe_get(job, "recommendation"),
        safe_get(job, "searchTerm"),
        safe_get(job, "source"),
        safe_get(job, "username"),
        safe_get(job, "role"),
        safe_get(job, "profilePath"),
        safe_get(job, "resumePath"),
    ]


def generate_timestamp() -> str:
    return datetime.now().strftime("%Y-%m-%d_%H-%M-%S")


def create_status_log(
    request: str,
    roles: list[str],
    filters: dict,
    results: dict,
    xlsx_path: str,
    status_dir: str | Path | None = None,
    timestamp: str | None = None,
) -> str:
    timestamp = timestamp or generate_timestamp()
    status_root = Path(status_dir) if status_dir else jobs_status_dir()
    status_root.mkdir(parents=True, exist_ok=True)
    status_path = status_root / f"{timestamp}.md"

    total_scraped = sum(r.get("scraped", 0) for r in results.values())
    total_filtered = sum(r.get("filtered", 0) for r in results.values())
    apply_count = sum(
        1
        for r in results.values()
        for j in r.get("jobs", [])
        if j.get("recommendation") == "Apply"
    )
    consider_count = sum(
        1
        for r in results.values()
        for j in r.get("jobs", [])
        if j.get("recommendation") == "Consider"
    )
    skip_count = sum(
        1
        for r in results.values()
        for j in r.get("jobs", [])
        if j.get("recommendation") == "Skip"
    )

    lines = [
        "# Job Scraping Status",
        "",
        f"**Generated**: {timestamp}",
        f'**User Request**: "{request}"',
        f"**Roles**: {', '.join(roles)}",
        f"**Filters**: {format_filters(filters)}",
        "",
        "## Results Summary",
        "",
        "| Source | Scraped | After Filter | Notes |",
        "|--------|---------|---------------|-------|",
    ]

    for source, data in results.items():
        lines.append(
            f"| {source.capitalize()} | {data.get('scraped', 0)} | {data.get('filtered', 0)} | {data.get('notes', '')} |"
        )

    lines.extend(
        [
            f"| **Total** | **{total_scraped}** | **{total_filtered}** | Target: {sum(r.get('target', 0) for r in results.values())} |",
            "",
            "## Files Generated",
            "",
            f"- {xlsx_path}",
            f"- {status_path}",
            "",
            "## Score Distribution",
            "",
            f"- Apply (>=70%): {apply_count}",
            f"- Consider (50-69%): {consider_count}",
            f"- Skip (<50%): {skip_count}",
            "",
        ]
    )

    status_path.write_text("\n".join(lines))
    return str(status_path)


def format_filters(filters: dict) -> str:
    parts = []
    if filters.get("location"):
        parts.append(f"location: {filters['location']}")
    if filters.get("work_mode"):
        parts.append(f"work_mode: {filters['work_mode']}")
    if filters.get("visa_status"):
        parts.append(f"visa: {filters['visa_status']}")
    if filters.get("score_threshold"):
        parts.append(f"score > {filters['score_threshold']}%")
    return ", ".join(parts) if parts else "default"


def consolidate_jobs(job_files: list[str]) -> list[dict]:
    jobs: list[dict] = []
    for filepath in job_files:
        path = Path(filepath)
        if not path.exists():
            continue
        try:
            loaded = json.loads(path.read_text())
        except json.JSONDecodeError:
            continue
        if isinstance(loaded, list):
            jobs.extend(loaded)
        elif isinstance(loaded, dict):
            jobs.append(loaded)
    return jobs


if __name__ == "__main__":
    test_jobs = [
        {
            "id": "1",
            "externalJobId": "EXT001",
            "title": "AI/ML Engineer",
            "jobDescription": "Looking for ML engineer",
            "salary": {
                "min": 120000,
                "max": 180000,
                "currencyCode": "USD",
                "type": "yearly",
            },
            "jobTypes": ["Full-time"],
            "workType": "remote",
            "location": {
                "city": "Boston",
                "state": "MA",
                "country": "US",
                "remote": True,
            },
            "company": {"name": "TechCorp", "industry": "AI"},
            "apply": {"applyUrl": "https://example.com/apply"},
            "jobSource": {"sourceName": "Dice"},
            "score": 75,
            "recommendation": "Apply",
            "searchTerm": "AI Engineer",
            "source": "dice",
        }
    ]
    output = str(jobs_dir() / f"{sanitize_filename('test_jobs')}.xlsx")
    print(jobs_to_csv(test_jobs, output))
