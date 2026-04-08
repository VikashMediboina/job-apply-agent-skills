#!/usr/bin/env python3
"""
extract_resume.py - Extract text from PDF or DOCX resume files.

Usage:
    python3 extract_resume.py <resume_path> [--output <output_path>]

Supports:
    - .pdf  (via pdfplumber, fallback to pdfminer, fallback to PyPDF2)
    - .docx (via python-docx)
    - .txt  (passthrough)

Output:
    Writes extracted markdown text to stdout or --output file.
    Preserves section structure where possible.
"""

import argparse
import os
import sys
import re


def extract_pdf(path: str) -> str:
    """Extract text from PDF using best available library."""
    text = ""

    # Try pdfplumber first (best layout preservation)
    try:
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n\n"
        if text.strip():
            return text.strip()
    except Exception:
        pass

    # Fallback to pdfminer (better for complex layouts)
    try:
        from pdfminer.high_level import extract_text as pdfminer_extract

        text = pdfminer_extract(path)
        if text.strip():
            return text.strip()
    except Exception:
        pass

    # Fallback to PyPDF2
    try:
        import PyPDF2

        with open(path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n\n"
        if text.strip():
            return text.strip()
    except Exception:
        pass

    return ""


def extract_docx(path: str) -> str:
    """Extract text from DOCX preserving paragraph structure."""
    import docx

    doc = docx.Document(path)
    lines = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            lines.append("")
            continue

        # Detect heading styles and convert to markdown headers
        style_name = (para.style.name or "").lower()
        if "heading 1" in style_name or "title" in style_name:
            lines.append(f"# {text}")
        elif "heading 2" in style_name:
            lines.append(f"## {text}")
        elif "heading 3" in style_name:
            lines.append(f"### {text}")
        elif "heading" in style_name:
            lines.append(f"#### {text}")
        else:
            # Check for bold runs that look like section headers
            if para.runs and all(r.bold for r in para.runs if r.text.strip()):
                if len(text) < 60:
                    lines.append(f"## {text}")
                else:
                    lines.append(text)
            else:
                lines.append(text)

    # Also extract from tables (common in resumes)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
        lines.append("")

    return "\n".join(lines).strip()


def extract_txt(path: str) -> str:
    """Read plain text file."""
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read().strip()


def normalize_text(text: str) -> str:
    """Clean up extracted text for downstream processing."""
    # Collapse excessive blank lines
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    # Remove trailing whitespace per line
    text = "\n".join(line.rstrip() for line in text.splitlines())
    # Strip leading/trailing
    return text.strip()


def detect_sections(text: str) -> dict:
    """
    Detect common resume sections from raw text.
    Returns dict of {section_name: section_content}.
    """
    # Common resume section headers (case-insensitive)
    section_patterns = [
        r"(?:professional\s+)?summary|objective|about\s*me|profile",
        r"(?:work\s+)?experience|employment\s+history|professional\s+experience",
        r"education|academic|qualifications",
        r"(?:technical\s+)?skills|competencies|technologies|expertise",
        r"projects|portfolio|key\s+projects",
        r"certifications?|licenses?",
        r"awards?|honors?|achievements?",
        r"publications?|papers?",
        r"volunteer|community",
        r"languages?",
        r"interests?|hobbies",
        r"references?",
        r"contact(?:\s+info(?:rmation)?)?",
    ]

    combined_pattern = "|".join(f"({p})" for p in section_patterns)
    header_re = re.compile(
        r"^(?:#{1,4}\s+)?(?:" + combined_pattern + r")\s*:?\s*$",
        re.IGNORECASE | re.MULTILINE,
    )

    sections = {}
    matches = list(header_re.finditer(text))

    if not matches:
        # No clear sections detected, return entire text as "full_resume"
        return {"full_resume": text}

    # Extract text before first section as "header" (name, contact)
    if matches[0].start() > 0:
        header_text = text[: matches[0].start()].strip()
        if header_text:
            sections["contact_header"] = header_text

    # Extract each section
    for i, match in enumerate(matches):
        section_name = (
            match.group(0).strip().strip("#").strip().strip(":").strip().lower()
        )
        start = match.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        section_content = text[start:end].strip()
        sections[section_name] = section_content

    return sections


def main():
    parser = argparse.ArgumentParser(
        description="Extract text from resume PDF/DOCX/TXT"
    )
    parser.add_argument("resume_path", help="Path to resume file")
    parser.add_argument("--output", "-o", help="Output file path (default: stdout)")
    parser.add_argument(
        "--sections",
        "-s",
        action="store_true",
        help="Also output detected sections as JSON",
    )
    args = parser.parse_args()

    path = os.path.expanduser(args.resume_path)
    if not os.path.isfile(path):
        print(f"Error: File not found: {path}", file=sys.stderr)
        sys.exit(1)

    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        raw_text = extract_pdf(path)
    elif ext in (".docx", ".doc"):
        raw_text = extract_docx(path)
    elif ext == ".txt":
        raw_text = extract_txt(path)
    else:
        print(f"Error: Unsupported file type: {ext}", file=sys.stderr)
        print("Supported: .pdf, .docx, .txt", file=sys.stderr)
        sys.exit(1)

    if not raw_text:
        print(f"Error: Could not extract text from: {path}", file=sys.stderr)
        sys.exit(1)

    cleaned = normalize_text(raw_text)

    if args.sections:
        import json

        sections = detect_sections(cleaned)
        output = json.dumps(
            {"raw_text": cleaned, "sections": sections},
            indent=2,
            ensure_ascii=False,
        )
    else:
        output = cleaned

    if args.output:
        os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)
        with open(args.output, "w", encoding="utf-8") as f:
            f.write(output)
        print(f"Extracted to: {args.output}", file=sys.stderr)
    else:
        print(output)


if __name__ == "__main__":
    main()
